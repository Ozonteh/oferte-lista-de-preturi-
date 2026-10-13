"""Importă o listă de prețuri în formatul vechi (Google Sheets exportat .xlsx)
într-un folder de furnizor: produse.csv + imagini/.

Formatul vechi are un rând de cap de tabel cu coloanele:
    Denumire | Poze | Link la sait | Preț | Preț Promo | Oferta
Linkurile stau ca hyperlink pe celule, iar pozele sunt plutitoare peste
coloana „Poze”. O poză se aplică tuturor produselor de la rândul ei în jos,
până la următoarea poză.

Citirea se oprește după 20 de rânduri goale la rând, ca să nu prindă datele
interne puse mai jos în foaie (prețuri de achiziție, containere).

Utilizare:
    python scripts/importa_lista_veche.py lista.xlsx furnizori/<furnizor>

După import, completați coloana `categorie` și corectați denumirile în
produse.csv, apoi rulați scripts/genereaza_lista.py.
"""

import csv
import io
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl
from PIL import Image

CAP_TABEL = {
    "denumire": "denumire",
    "poze": "poze",
    "link la sait": "link",
    "link la site": "link",
    "preț": "pret",
    "pret": "pret",
    "preț promo": "promo",
    "pret promo": "promo",
    "oferta": "oferta",
    "ofertă": "oferta",
}
RANDURI_GOALE_STOP = 20
LATURA_MAX_IMAGINE = 400
COLOANE_CSV = [
    "categorie", "denumire", "link_produs", "pret_eur", "pret_promo_eur",
    "nota", "link_oferta", "imagine",
]


def normalizeaza(text):
    return re.sub(r"\s+", " ", str(text or "")).strip()


def slug(text):
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-") or "produs"


def gaseste_cap_tabel(ws):
    for row in ws.iter_rows(max_row=50):
        coloane = {}
        for cell in row:
            cheie = CAP_TABEL.get(normalizeaza(cell.value).lower())
            if cheie:
                coloane[cheie] = cell.column
        if "denumire" in coloane and "pret" in coloane:
            return row[0].row, coloane
    raise SystemExit("Nu am găsit rândul de cap de tabel (Denumire / Preț).")


def numar_sau_text(valoare):
    if isinstance(valoare, (int, float)):
        return f"{valoare:g}"
    text = normalizeaza(valoare)
    return "" if text in ("-", "link") else text


def main(cale_xlsx, folder_furnizor):
    folder = Path(folder_furnizor)
    (folder / "imagini").mkdir(parents=True, exist_ok=True)
    ws = openpyxl.load_workbook(cale_xlsx).active
    rand_cap, col = gaseste_cap_tabel(ws)

    # Poze: rândul de ancorare -> imagine. Pozele de deasupra capului de tabel
    # (de obicei sigla furnizorului) se salvează separat ca logo.png.
    poze = []
    for img in ws._images:
        rand = img.anchor._from.row + 1
        imagine = Image.open(io.BytesIO(img._data())).convert("RGB")
        if rand < rand_cap:
            imagine.save(folder / "logo.png")
        else:
            poze.append((rand, imagine))
    poze.sort(key=lambda p: p[0])

    produse, goale = [], 0
    for rand in range(rand_cap + 1, ws.max_row + 1):
        celula_nume = ws.cell(rand, col["denumire"])
        denumire = normalizeaza(celula_nume.value)
        if not denumire:
            goale += 1
            if goale >= RANDURI_GOALE_STOP:
                break
            continue
        goale = 0

        def link(cheie):
            if cheie not in col:
                return ""
            celula = ws.cell(rand, col[cheie])
            return celula.hyperlink.target if celula.hyperlink else ""

        pret = numar_sau_text(ws.cell(rand, col["pret"]).value)
        promo = numar_sau_text(ws.cell(rand, col["promo"]).value) if "promo" in col else ""
        nota = ""
        if promo and not re.fullmatch(r"[\d.]+", promo):
            nota, promo = promo, ""
        produse.append({
            "rand": rand,
            "categorie": "",
            "denumire": denumire,
            "link_produs": link("link"),
            "pret_eur": pret if re.fullmatch(r"[\d.]+", pret) else "",
            "pret_promo_eur": promo,
            "nota": nota,
            "link_oferta": link("oferta"),
            "imagine": "",
        })

    # Fiecare produs ia ultima poză ancorată la sau deasupra rândului său.
    salvate = {}
    for produs in produse:
        sursa = [p for p in poze if p[0] <= produs["rand"]]
        if not sursa:
            continue
        rand_poza, imagine = sursa[-1]
        if rand_poza not in salvate:
            nume = f"{slug(produs['denumire'])}.jpg"
            imagine.thumbnail((LATURA_MAX_IMAGINE, LATURA_MAX_IMAGINE))
            imagine.save(folder / "imagini" / nume, quality=90)
            salvate[rand_poza] = nume
        produs["imagine"] = salvate[rand_poza]

    with open(folder / "produse.csv", "w", newline="", encoding="utf-8-sig") as f:
        scriitor = csv.DictWriter(f, fieldnames=COLOANE_CSV, extrasaction="ignore")
        scriitor.writeheader()
        scriitor.writerows(produse)

    print(f"{len(produse)} produse și {len(salvate)} poze -> {folder}")
    fara_pret = [p["denumire"] for p in produse if not p["pret_eur"]]
    if fara_pret:
        print("Fără preț (apar „La cerere”):", "; ".join(fara_pret))


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit(__doc__)
    main(sys.argv[1], sys.argv[2])
